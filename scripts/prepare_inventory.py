#!/usr/bin/env python3
"""Prepare a review queue from the PDR Kampüs inventory; never publish unverified files.

Usage: python scripts/prepare_inventory.py inventory.xlsx --out-dir inventory-review
Requires: pip install openpyxl (only for reading the input workbook).
"""
import argparse
import csv
import json
import re
import unicodedata
from collections import Counter, defaultdict
from pathlib import Path
from urllib.parse import unquote, urlsplit, urlunsplit

HEADERS = ('Başlık', 'Kategori', 'Kademe / Sınıf', 'Hedef Kitle', 'Materyal Türü', 'Öncelik', 'Kaynak Kurum', 'Kaynak URL', 'Not')
DIRECT_EXTENSIONS = ('.pdf', '.xlsx', '.xls', '.docx', '.doc', '.pptx', '.ppt', '.zip')
OUTPUT_HEADERS = ('row', 'status', 'title', 'category', 'level', 'audience', 'material_type', 'priority', 'institution', 'source_url', 'candidate_file', 'matched_catalog_id', 'match_reason', 'duplicate_rows', 'notes')
PRIORITY = {'Çok yüksek': 0, 'Yüksek': 1, 'Orta': 2, 'Düşük': 3}


def norm(value):
    value = unicodedata.normalize('NFKD', str(value or '').casefold())
    return re.sub(r'[^a-z0-9]+', ' ', ''.join(c for c in value if not unicodedata.combining(c)).replace('ı', 'i')).strip()


def canonical_url(value):
    value = str(value or '').strip()
    if not value:
        return ''
    parts = urlsplit(value)
    if parts.scheme not in ('http', 'https') or not parts.netloc:
        return ''
    return urlunsplit((parts.scheme.lower(), parts.netloc.lower(), parts.path, '', ''))


def file_url(url):
    return url if unquote(urlsplit(url).path).lower().endswith(DIRECT_EXTENSIONS) else ''


def load_inventory(path):
    if path.suffix.lower() == '.csv':
        with path.open(encoding='utf-8-sig', newline='') as handle:
            yield from csv.DictReader(handle)
        return
    if path.suffix.lower() != '.xlsx':
        raise ValueError('Only .xlsx or .csv inventory files are supported')
    try:
        from openpyxl import load_workbook
    except ImportError as error:
        raise SystemExit('XLSX reading requires: pip install openpyxl') from error
    book = load_workbook(path, read_only=True, data_only=True)
    if 'Kaynak Envanteri' not in book:
        raise ValueError('Sheet "Kaynak Envanteri" is missing')
    sheet = book['Kaynak Envanteri']
    rows = sheet.iter_rows(values_only=True)
    headers = [str(x or '').strip() for x in next(rows)]
    for values in rows:
        yield dict(zip(headers, values))
    book.close()


def prepare(rows, library, forms):
    records = []
    catalog = library + forms
    by_file = {canonical_url(x.get('file')): x['id'] for x in catalog if x.get('file')}
    by_title = defaultdict(list)
    for x in catalog:
        by_title[norm(x.get('title'))].append(x['id'])
    for row_number, source in enumerate(rows, start=2):
        if not str(source.get('Başlık') or '').strip():
            continue
        missing = [name for name in HEADERS if name not in source]
        if missing:
            raise ValueError('Inventory missing columns: ' + ', '.join(missing))
        url = canonical_url(source.get('Kaynak URL'))
        direct = file_url(url)
        title = str(source['Başlık']).strip()
        matches = by_title[norm(title)]
        if direct and direct in by_file:
            status, matched, reason = 'published', by_file[direct], 'same_file_url'
        elif len(matches) == 1:
            status, matched, reason = 'review', matches[0], 'same_title_check_file'
        elif direct:
            status, matched, reason = 'file_candidate', '', 'verify_file_and_metadata'
        else:
            status, matched, reason = 'review', '', 'find_specific_file_on_source_page' if url else 'missing_source_url'
        records.append({
            'row': row_number, 'status': status, 'title': title,
            'category': str(source.get('Kategori') or ''), 'level': str(source.get('Kademe / Sınıf') or ''),
            'audience': str(source.get('Hedef Kitle') or ''), 'material_type': str(source.get('Materyal Türü') or ''),
            'priority': str(source.get('Öncelik') or ''), 'institution': str(source.get('Kaynak Kurum') or ''),
            'source_url': url, 'candidate_file': direct, 'matched_catalog_id': matched,
            'match_reason': reason, 'duplicate_rows': '', 'notes': str(source.get('Not') or '')
        })
    same_title = defaultdict(list)
    same_file = defaultdict(list)
    for record in records:
        same_title[norm(record['title'])].append(record['row'])
        if record['candidate_file']:
            same_file[record['candidate_file']].append(record['row'])
    for record in records:
        siblings = set(same_title[norm(record['title'])] + same_file.get(record['candidate_file'], [])) - {record['row']}
        record['duplicate_rows'] = ','.join(map(str, sorted(siblings)))
        if siblings and record['status'] == 'file_candidate':
            record['status'] = 'review'
            record['match_reason'] = 'duplicate_candidate_check_file'
    return records


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('inventory', type=Path)
    parser.add_argument('--out-dir', type=Path, default=Path('inventory-review'))
    args = parser.parse_args()
    root = Path(__file__).resolve().parents[1]
    library = json.loads((root / 'data/library.json').read_text(encoding='utf-8'))
    forms = json.loads((root / 'data/forms.json').read_text(encoding='utf-8'))
    records = prepare(load_inventory(args.inventory), library, forms)
    args.out_dir.mkdir(parents=True, exist_ok=True)
    with (args.out_dir / 'review-queue.csv').open('w', newline='', encoding='utf-8-sig') as handle:
        writer = csv.DictWriter(handle, fieldnames=OUTPUT_HEADERS)
        writer.writeheader()
        writer.writerows(sorted(records, key=lambda r: (0 if r['status'] == 'file_candidate' else 1 if r['status'] == 'review' else 2,
                                                         PRIORITY.get(r['priority'], 4), r['row'])))
    pages = defaultdict(list)
    for record in records:
        if record['status'] == 'review' and record['source_url'] and not record['candidate_file']:
            pages[record['source_url']].append(record)
    with (args.out_dir / 'source-pages.csv').open('w', newline='', encoding='utf-8-sig') as handle:
        writer = csv.writer(handle)
        writer.writerow(('source_url', 'inventory_titles', 'high_priority', 'already_cataloged_files', 'example_titles'))
        catalog_pages = Counter(canonical_url(x.get('sourcePage')) for x in library + forms)
        for url, group in sorted(pages.items(), key=lambda p: (-sum(x['priority'] == 'Çok yüksek' for x in p[1]), -len(p[1]), p[0])):
            writer.writerow((url, len(group), sum(x['priority'] == 'Çok yüksek' for x in group),
                             catalog_pages[url], ' | '.join(x['title'] for x in group[:3])))
    counts = Counter(record['status'] for record in records)
    summary = {'inventory_rows': len(records), 'catalog_library': len(library), 'catalog_forms': len(forms),
               'published_by_url': counts['published'], 'file_candidates': counts['file_candidate'],
               'manual_review': counts['review'],
               'duplicate_rows': sum(bool(r['duplicate_rows']) for r in records),
               'unique_source_pages': len({r['source_url'] for r in records if r['source_url'] and not r['candidate_file']})}
    (args.out_dir / 'summary.json').write_text(json.dumps(summary, ensure_ascii=False, indent=2) + '\n', encoding='utf-8')
    print(json.dumps(summary, ensure_ascii=False, indent=2))
    print('Review queue:', args.out_dir / 'review-queue.csv')
    print('Grouped source pages:', args.out_dir / 'source-pages.csv')


if __name__ == '__main__':
    main()
