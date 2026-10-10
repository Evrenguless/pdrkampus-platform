"""Build faithful previews from public source files; never modify resource records."""
import time
import io, re, zipfile
import xml.etree.ElementTree as ET
import shutil
import argparse, concurrent.futures, hashlib, json, os, subprocess, tempfile, textwrap, urllib.request, urllib.parse
from pathlib import Path
from openpyxl import load_workbook
from PIL import Image, ImageDraw, ImageFont
from pypdf import PdfReader
class ConverterUnavailableError(RuntimeError):pass
ROOT=Path(__file__).resolve().parents[1]
FONT=os.environ.get('PREVIEW_FONT') or next((p for p in ['/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf','/System/Library/Fonts/Helvetica.ttc','C:/Windows/Fonts/arial.ttf'] if Path(p).exists()),'DejaVuSans.ttf')
def thumbnail(sheet,path):
    rows=sheet['rows'][:16];cols=min(max(map(len,rows),default=1),8)
    width=1200;cw=(width-70)//cols;rh=82
    canvas=Image.new('RGB',(width,64+len(rows)*rh),'white');d=ImageDraw.Draw(canvas)
    font=ImageFont.truetype(FONT,17);small=ImageFont.truetype(FONT,14)
    d.rectangle((0,0,width,58),fill='#e8f2ed');d.text((22,18),sheet['name'],font=font,fill='#183f31')
    covered=set();merges={}
    from openpyxl.utils.cell import range_boundaries
    for value in sheet.get('merges',[]):
        sc,sr,ec,er=range_boundaries(value);sr-=1;sc-=1;er=min(er,len(rows));ec=min(ec,cols)
        if sr>=len(rows) or sc>=cols:continue
        merges[sr,sc]=(er-sr,ec-sc)
        covered.update((y,x) for y in range(sr,er) for x in range(sc,ec) if (y,x)!=(sr,sc))
    for y,row in enumerate(rows):
        d.text((7,64+y*rh+10),str(y+1),font=small,fill='#69766f')
        for x in range(cols):
            if (y,x) in covered:continue
            rowspan,colspan=merges.get((y,x),(1,1));left=45+x*cw;top=64+y*rh;right=left+cw*colspan;bottom=top+rh*rowspan
            d.rectangle((left,top,right,bottom),outline='#b7c8bf',fill='#f4f8f5' if y<2 else 'white')
            value=str(row[x] if x<len(row) else '')
            lines=[]
            for line in value.splitlines():lines.extend(textwrap.wrap(line,width=max(9,int((right-left-16)/9))) or [''])
            for n,line in enumerate(lines[:max(1,(bottom-top-12)//22)]):d.text((left+8,top+7+n*22),line,font=font,fill='#173728')
    canvas.save(path,'JPEG',quality=80,optimize=True)
def preview_workbook(path,data_only):
    try:
        return load_workbook(path,data_only=data_only)
    except ValueError as error:
        # Some official books contain an unsupported drawing font pitchFamily.
        # Remove drawing references in an in-memory preview copy only; keep cells,
        # formulas, cached results, merges and the original downloaded bytes intact.
        if str(error.__cause__)!='Max value is 52':raise
        cleaned=io.BytesIO()
        with zipfile.ZipFile(path) as source, zipfile.ZipFile(cleaned,'w') as target:
            for entry in source.infolist():
                content=source.read(entry.filename)
                if re.fullmatch(r'xl/worksheets/sheet\d+\.xml',entry.filename):
                    content=re.sub(rb'<drawing\b[^>]*/>',b'',content)
                if entry.filename.startswith('xl/worksheets/_rels/') and entry.filename.endswith('.rels'):
                    relations=ET.fromstring(content)
                    for relation in list(relations):
                        if relation.get('Type','').endswith('/drawing'):relations.remove(relation)
                    content=ET.tostring(relations)
                target.writestr(entry,content)
        cleaned.seek(0)
        return load_workbook(cleaned,data_only=data_only)

def build(row, *, root=None, source_body=None):
    root = Path(root) if root is not None else ROOT
    kind=row['fileType'].upper();ident=row['id'];assets=root/'assets/resource-previews';assets.mkdir(parents=True,exist_ok=True)
    with tempfile.TemporaryDirectory() as tmp:
        dest=Path(tmp)/('source.'+kind.lower())
        override=json.loads((root/'seo/link-overrides.json').read_text()).get('entries',{}).get(row['file'],{})
        url=override.get('replacement_url') if row['id'] in override.get('ids',[]) and override.get('replacement_url') else row['file']
        parts=urllib.parse.urlsplit(url)
        url=urllib.parse.urlunsplit((parts.scheme,parts.netloc,urllib.parse.quote(parts.path,safe='/%:@'),urllib.parse.quote(parts.query,safe='/?&=%:+;'),parts.fragment))
        request=urllib.request.Request(url,headers={'User-Agent':'Mozilla/5.0 (compatible; PDRKampus-DocumentPreview/1.1)'})
        if source_body is None:
            with urllib.request.urlopen(request,timeout=60) as response:data=response.read(100_000_001)
        else:
            data=source_body
        expected=row.get('contentSha256')
        if source_body is not None and expected and hashlib.sha256(data).hexdigest()!=expected:raise ValueError('Preview source hash mismatch')
        if len(data)>100_000_000:raise ValueError('Preview size limit')
        dest.write_bytes(data);result={'sourceSha256':hashlib.sha256(data).hexdigest()}
        if kind in ['DOC','DOCX','PPT','PPTX','XLS']:
            if not shutil.which('libreoffice'):raise ConverterUnavailableError('Install LibreOffice for Office document previews')
            profile=Path(tmp)/'profile';(profile/'user').mkdir(parents=True)
            (profile/'user/registrymodifications.xcu').write_text('<?xml version="1.0"?><oor:items xmlns:oor="http://openoffice.org/2001/registry"><item oor:path="/org.openoffice.Office.Common/Security/Scripting"><prop oor:name="MacroSecurityLevel" oor:op="fuse"><value>3</value></prop></item><item oor:path="/org.openoffice.Office.Common/Security/Scripting"><prop oor:name="DisableMacrosExecution" oor:op="fuse"><value>true</value></prop></item></oor:items>')
            subprocess.run(['libreoffice','-env:UserInstallation='+profile.as_uri(),'--headless','--convert-to','pdf','--outdir',tmp,str(dest)],check=True,capture_output=True,timeout=60)
            dest=Path(tmp)/'source.pdf'
            if not dest.exists():raise ValueError('Office conversion did not produce a PDF')
            result['sourceFormat']=kind
            kind='PDF'
        if kind=='PDF':
            reader=PdfReader(dest);result['pageCount']=len(reader.pages)
            prefix=assets/ident
            subprocess.run(['pdftoppm','-f','1','-l','2','-scale-to','1200','-jpeg','-jpegopt','quality=78',str(dest),str(prefix)],check=True,capture_output=True,timeout=90)
            result['previews']=['/'+str(p.relative_to(root)) for p in sorted(assets.glob(ident+'-*.jpg'))]
            result['thumbnail']=result['previews'][0]
        elif kind=='XLSX':
            book=preview_workbook(dest,data_only=True);formulas=preview_workbook(dest,data_only=False);sheets=[]
            for sheet in book:
                if sheet.sheet_state!='visible':continue
                values=[[str(c.value) if c.value is not None else ('Formül sonucu özgün Excel dosyasında görüntülenir' if formulas[sheet.title].cell(c.row,c.column).data_type=='f' else '') for c in cells] for cells in sheet.iter_rows(max_row=min(sheet.max_row,180),max_col=min(sheet.max_column,40))]
                while values and not any(values[-1]):values.pop()
                if not values:continue
                used=max((i+1 for row in values for i,v in enumerate(row) if v),default=1)
                sheets.append({'name':sheet.title,'rows':[row[:used] for row in values],'merges':[str(v) for v in sheet.merged_cells.ranges],'rowCount':sheet.max_row,'columnCount':sheet.max_column})
            if not sheets:raise ValueError('No visible sheet')
            result.update(sheets=sheets,sheetCount=len(sheets));chosen=next((s for s in sheets if s['name']=='EYLÜL'),sheets[0]);cover=assets/(ident+'-sheet.jpg');thumbnail(chosen,cover);result['thumbnail']='/'+str(cover.relative_to(root))
        return result

def main():
    parser=argparse.ArgumentParser();parser.add_argument('--limit',type=int,default=100);args=parser.parse_args()
    if not 0<args.limit<=5000:parser.error('--limit must be between 1 and 5000')
    manifest=ROOT/'data/resource-previews.json';previews=json.loads(manifest.read_text())
    state_file=ROOT/'data/preview-build-state.json';state=json.loads(state_file.read_text()) if state_file.exists() else {}
    rows=[]
    for name in ['forms','library','collected-resources','library-reviewed-batch-20261010']:rows+=json.loads((ROOT/f'data/{name}.json').read_text())
    # The library sorts titles; prioritize visible cards and the explicitly reported workbook.
    rows.sort(key=lambda r:(r['id']!='collected-64cf8183f040bf20432c3eea',r['id'] not in { 'tuba-teknoloji-bagimliligi-raporu', 'meb-okuloncesi-uyum-etkinlikleri-2025', 'meb-aile-ici-iletisim-sunumu-2023', 'dergipdr-empati-gelistirme-rehberi', 'yalvac-ram-verimli-ders-calisma-etkinlikleri', 'iznik-ortaokul-uyum-rehberi' },r['title']))
    by_file={r['file']:previews[r['id']] for r in rows if r['id'] in previews and previews[r['id']].get('thumbnail')}
    for row in rows:
        if row['fileType'].upper() in ['PNG','JPG','JPEG']:
            by_file.setdefault(row['file'],{'thumbnail':row['file'],'previews':[row['file']]})
    pending=[];seen=set()
    for row in rows:
        if row['file'] in by_file:previews[row['id']]=by_file[row['file']];state.pop(row['id'],None);continue
        if row['fileType'].upper() not in ['PDF','XLSX','DOC','DOCX','PPT','PPTX','XLS'] or row['file'] in seen:continue
        if state.get(row['id'],{}).get('message') and state.get(row['id'],{}).get('retryAfter',0)>time.time() and state.get(row['id'],{}).get('errorType') not in (['UnicodeEncodeError','ConverterUnavailableError'] if shutil.which('libreoffice') else ['UnicodeEncodeError']):continue
        seen.add(row['file']);pending.append(row)
    successes=0
    with concurrent.futures.ThreadPoolExecutor(max_workers=6) as pool:
        futures={pool.submit(build,row):row for row in pending[:args.limit]}
        for future in concurrent.futures.as_completed(futures):
            row=futures[future]
            try:result=future.result();by_file[row['file']]=result;successes+=1;state.pop(row['id'],None);print('Preview ready:',row['id'],flush=True)
            except Exception as error:
                state[row['id']]={'retryAfter':int(time.time()+86400),'errorType':type(error).__name__,'message':str(error)[:500]}
                print('Preview deferred:',row['id'],type(error).__name__,flush=True)
    for row in rows:
        if row['file'] in by_file:previews[row['id']]=by_file[row['file']]
    manifest.write_text(json.dumps(previews,ensure_ascii=False,separators=(',',':'))+'\n')
    state_file.write_text(json.dumps(state,indent=2)+'\n')
    missing=[r for r in rows if r['id'] not in previews and r['fileType'].upper() in ['PDF','XLSX','DOC','DOCX','PPT','PPTX','XLS']]
    print(json.dumps({'missing_document_records':len({r['id'] for r in missing}),'deferred_records':len({r['id'] for r in missing if state.get(r['id'],{}).get('retryAfter',0)>time.time()}),'generated':successes,'records_with_preview':len(previews),'remaining_files':max(0,len(pending)-successes)}))
if __name__=='__main__':main()
